#!/usr/bin/env python3
"""Docker helper probing fixtures through the MCP tools agents use.

This file is executed inside the benchmark environment image by
``test_generated_fixture_mcp_docker.py``. Each invocation uses one MCP
server's own virtual environment to avoid cross-server module-name collisions.
"""

from __future__ import annotations

import json
import os
import re
import shlex
import sys
from pathlib import Path


# The probe file is mounted at /fixture_probe.py, while docker exec selects the
# relevant MCP server directory as cwd. Python otherwise adds only the script's
# directory (/) to sys.path.
sys.path.insert(0, os.getcwd())


MARKER = "ARCHIPELAGO_MCP_WRITE_PROBE"


def _args() -> tuple[str, list[str]]:
    if len(sys.argv) != 3:
        raise SystemExit("usage: generated_fixture_mcp_probe.py KIND JSON_PATHS")
    paths = json.loads(sys.argv[2])
    if not isinstance(paths, list) or not all(isinstance(path, str) for path in paths):
        raise ValueError("JSON_PATHS must be an array of strings")
    return sys.argv[1], paths


def probe_filesystem(paths: list[str]) -> dict[str, object]:
    from tools.read_text_file import read_text_file

    sizes = {}
    for path in paths:
        content = read_text_file.__wrapped__(path)
        if not content:
            raise AssertionError(f"Filesystem MCP returned empty content for {path}")
        sizes[path] = len(content)
    return {"read": sizes}


def probe_code_modify(paths: list[str]) -> dict[str, object]:
    from models.code_exec import CodeExecRequest
    from tools.code_exec import code_exec, verify_sandbox_available

    verify_sandbox_available()
    commands = [
        f"printf '\\n{MARKER}\\n' >> {shlex.quote(path.lstrip('/'))}" for path in paths
    ]
    result = code_exec.__wrapped__(CodeExecRequest(code=" && ".join(commands)))
    if not result.success:
        raise AssertionError(f"Code-execution MCP could not modify fixture: {result.output}")
    return {"modified": paths}


def probe_filesystem_marker(paths: list[str]) -> dict[str, object]:
    from tools.read_text_file import read_text_file

    for path in paths:
        content = read_text_file.__wrapped__(path)
        if MARKER not in content:
            raise AssertionError(f"Filesystem MCP did not observe code-execution edit: {path}")
    return {"verified_modified": paths}


def probe_spreadsheets(paths: list[str]) -> dict[str, object]:
    from openpyxl import load_workbook
    from tools.edit_spreadsheet import EditSpreadsheetInput, edit_spreadsheet
    from tools.list_tabs_in_spreadsheet import list_tabs_in_spreadsheet
    from tools.read_tab import ReadTabInput, read_tab

    verified = []
    for path in paths:
        tabs = list_tabs_in_spreadsheet.__wrapped__(path)
        if "Worksheet" not in tabs and "Tab" not in tabs and "sheet" not in tabs.lower():
            raise AssertionError(f"Sheets MCP could not list workbook tabs for {path}: {tabs}")
        physical = Path("/filesystem") / path.lstrip("/")
        workbook = load_workbook(physical, read_only=True)
        try:
            sheet = workbook.sheetnames[0]
        finally:
            workbook.close()
        edit_result = edit_spreadsheet.__wrapped__(
            EditSpreadsheetInput(
                file_path=path,
                operations=[
                    {"type": "set_cell", "sheet": sheet, "cell": "ZZ100", "value": MARKER}
                ],
            )
        )
        if "success" not in edit_result.lower():
            raise AssertionError(f"Sheets MCP edit failed for {path}: {edit_result}")
        value = read_tab.__wrapped__(
            ReadTabInput(file_path=path, tab_index=0, cell_range="ZZ100")
        )
        if MARKER not in value:
            raise AssertionError(f"Sheets MCP did not read its edit for {path}: {value}")
        verified.append(path)
    return {"read_and_modified": verified}


def probe_mail(_: list[str]) -> dict[str, object]:
    from models.mail import SendMailInput
    from tools.list_mails import list_mails
    from tools.read_mail import read_mail
    from tools.send_mail import send_mail

    before = list_mails.__wrapped__()
    ids = re.findall(r"Mail ID: (\S+)", before)
    if not ids:
        raise AssertionError(f"Mail MCP returned no populated email: {before}")
    original = read_mail.__wrapped__(ids[0])
    if "not found" in original.lower() or "failed" in original.lower():
        raise AssertionError(f"Mail MCP could not read {ids[0]}: {original}")
    sent = send_mail.__wrapped__(
        SendMailInput(
            from_email="probe@benchmark.local",
            to_email="receiver@benchmark.local",
            subject=MARKER,
            body=MARKER,
        )
    )
    if "success" not in sent.lower() and "sent" not in sent.lower():
        raise AssertionError(f"Mail MCP write failed: {sent}")
    after = list_mails.__wrapped__()
    if MARKER not in after:
        raise AssertionError("Mail MCP did not list the newly sent probe email")
    return {"read_mail_id": ids[0], "modified": True}


def probe_calendar(_: list[str]) -> dict[str, object]:
    from models.calendar import ListEventsRequest, ReadEventRequest, UpdateEventRequest
    from tools.list_events import list_events_sync
    from tools.read_event import read_event
    from tools.update_event import update_event

    events = list_events_sync(ListEventsRequest())
    if not events.events:
        raise AssertionError("Calendar MCP returned no populated event")
    event_id = events.events[0].id
    before = read_event.__wrapped__(ReadEventRequest(event_id=event_id))
    if not before:
        raise AssertionError(f"Calendar MCP could not read {event_id}")
    updated = update_event.__wrapped__(
        UpdateEventRequest(event_id=event_id, summary=MARKER)
    )
    if "success" not in updated.lower() and "updated" not in updated.lower():
        raise AssertionError(f"Calendar MCP update failed: {updated}")
    after = read_event.__wrapped__(ReadEventRequest(event_id=event_id))
    if MARKER not in after:
        raise AssertionError("Calendar MCP did not read its updated event")
    return {"read_event_id": event_id, "modified": True}


def probe_chat(_: list[str]) -> dict[str, object]:
    from models.requests import GetChannelHistoryRequest, ListChannelsRequest, PostMessageRequest
    from tools.get_channel_history import get_channel_history
    from tools.list_channels import list_channels
    from tools.post_message import post_message

    channels = list_channels.__wrapped__(ListChannelsRequest())
    if not channels.groups:
        raise AssertionError("Chat MCP returned no populated channel")
    channel_id = channels.groups[0].id
    before = get_channel_history.__wrapped__(GetChannelHistoryRequest(channel_id=channel_id))
    if not before.messages:
        raise AssertionError(f"Chat MCP returned empty history for {channel_id}")
    post_message.__wrapped__(PostMessageRequest(channel_id=channel_id, message=MARKER))
    after = get_channel_history.__wrapped__(GetChannelHistoryRequest(channel_id=channel_id))
    if not any(message.text == MARKER for message in after.messages):
        raise AssertionError("Chat MCP did not return its newly posted message")
    return {"channel_id": channel_id, "initial_messages": len(before.messages), "modified": True}


def main() -> None:
    kind, paths = _args()
    probes = {
        "filesystem": probe_filesystem,
        "code_modify": probe_code_modify,
        "filesystem_marker": probe_filesystem_marker,
        "spreadsheets": probe_spreadsheets,
        "mail": probe_mail,
        "calendar": probe_calendar,
        "chat": probe_chat,
    }
    if kind not in probes:
        raise ValueError(f"unknown probe kind: {kind}")
    print(json.dumps({"kind": kind, "ok": True, **probes[kind](paths)}))


if __name__ == "__main__":
    main()
