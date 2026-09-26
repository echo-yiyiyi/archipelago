"""Extract font, color, layout, and shape metadata from a PPTX file.

This is a domain-specific utility called only by the pptx_style_verifier eval,
NOT registered in the TRANSFORMATION_REGISTRY.
"""

from io import BytesIO
from typing import Any

from pptx import Presentation
from pptx.util import Emu


def _emu_to_pt(emu_val: int | Emu | None) -> float | None:
    """Convert EMU to points. Returns None if input is None."""
    if emu_val is None:
        return None
    return round(int(emu_val) / 12700, 1)


def _rgb_to_hex(rgb: Any) -> str | None:
    """Convert an RGBColor to a hex string like '#RRGGBB'. Returns None if not set."""
    if rgb is None:
        return None
    try:
        return f"#{rgb}"
    except Exception:
        return None


def _extract_run_style(run: Any) -> dict[str, Any]:
    """Extract style properties from a single text run."""
    font = run.font
    try:
        color_rgb = _rgb_to_hex(font.color.rgb) if font.color else None
    except AttributeError:
        color_rgb = None
    return {
        "text": run.text,
        "font_name": font.name,  # None means inherited from slide master
        "font_size_pt": _emu_to_pt(font.size),
        "font_color_rgb": color_rgb,
        "bold": font.bold,  # None means inherited
        "italic": font.italic,  # None means inherited
    }


def _extract_shape_metadata(shape: Any) -> dict[str, Any]:
    """Extract metadata from a single shape."""
    result: dict[str, Any] = {
        "shape_name": shape.name,
        "shape_type": str(shape.shape_type) if shape.shape_type else None,
        "left": _emu_to_pt(shape.left),
        "top": _emu_to_pt(shape.top),
        "width": _emu_to_pt(shape.width),
        "height": _emu_to_pt(shape.height),
    }

    if shape.has_text_frame:
        paragraphs = []
        for para in shape.text_frame.paragraphs:
            runs = [_extract_run_style(run) for run in para.runs]
            paragraphs.append(
                {
                    "alignment": str(para.alignment) if para.alignment else None,
                    "level": para.level,
                    "text_runs": runs,
                }
            )
        result["paragraphs"] = paragraphs

    return result


def pptx_to_style_metadata(file_bytes: bytes, file_name: str) -> dict[str, Any]:
    """Extract style metadata from a PPTX file.

    Returns:
        {
            "file_name": str,
            "slide_width_pt": float,
            "slide_height_pt": float,
            "slide_count": int,
            "slides": [
                {
                    "index": int,
                    "title": str | None,
                    "layout_name": str,
                    "shapes": [...]
                }
            ]
        }
    """
    prs = Presentation(BytesIO(file_bytes))

    slide_width_pt = _emu_to_pt(prs.slide_width)
    slide_height_pt = _emu_to_pt(prs.slide_height)

    slides_meta: list[dict[str, Any]] = []
    for idx, slide in enumerate(prs.slides):
        # Extract title from the title placeholder if present
        title = None
        if slide.shapes.title and slide.shapes.title.has_text_frame:
            title = slide.shapes.title.text_frame.text

        layout_name = slide.slide_layout.name if slide.slide_layout else None

        shapes = [_extract_shape_metadata(s) for s in slide.shapes]

        slides_meta.append(
            {
                "index": idx,
                "title": title,
                "layout_name": layout_name,
                "shapes": shapes,
            }
        )

    return {
        "file_name": file_name,
        "slide_width_pt": slide_width_pt,
        "slide_height_pt": slide_height_pt,
        "slide_count": len(slides_meta),
        "slides": slides_meta,
    }
